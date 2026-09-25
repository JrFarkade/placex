import os
import sys
import json
import docx

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"C:\Users\JrFar\.gemini\antigravity\scratch\placex\backend\PlaceX_Roadmap_Corrected_Reference.docx"
OUTPUT_JSON_PATH = r"C:\Users\JrFar\.gemini\antigravity\scratch\placex\backend\app\learning_engine\data\roadmap_docx_data.json"

BRANCHES_MAPPING = [
    ("Data Science", ["Beginner", "Intermediate", "Advanced"]),
    ("AI/ML Engineering", ["Beginner", "Intermediate", "Advanced"]),
    ("Cybersecurity", ["Beginner", "Intermediate", "Advanced"]),
    ("Computer Science / Software Development", ["Beginner", "Intermediate", "Advanced"])
]

def parse_docx():
    if not os.path.exists(DOCX_PATH):
        print(f"Error: DOCX file not found at {DOCX_PATH}")
        sys.exit(1)

    doc = docx.Document(DOCX_PATH)
    print(f"Total tables in document: {len(doc.tables)}")
    
    if len(doc.tables) != 12:
        print(f"Warning: Expected 12 tables, found {len(doc.tables)}")

    data_store = {}
    table_idx = 0

    for branch_name, levels in BRANCHES_MAPPING:
        data_store[branch_name] = {}
        for level in levels:
            if table_idx >= len(doc.tables):
                break
            
            table = doc.tables[table_idx]
            headers = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells]
            
            weeks_list = []
            for r_idx in range(1, len(table.rows)):
                cells = table.rows[r_idx].cells
                row_data = {}
                for h, cell in zip(headers, cells):
                    row_data[h] = cell.text.strip()
                
                try:
                    wk_num = int(row_data.get('Wk', r_idx))
                except ValueError:
                    wk_num = r_idx
                
                week_entry = {
                    "week": wk_num,
                    "topic": row_data.get('Topic', ''),
                    "difficulty": row_data.get('Difficulty', ''),
                    "priority": row_data.get('Priority', ''),
                    "prerequisites": row_data.get('Prerequisites', ''),
                    "completion_criteria": row_data.get('Completion Criteria', '')
                }
                weeks_list.append(week_entry)
            
            data_store[branch_name][level] = weeks_list
            print(f"Parsed {branch_name} -> {level}: {len(weeks_list)} weeks")
            table_idx += 1

    os.makedirs(os.path.dirname(OUTPUT_JSON_PATH), exist_ok=True)
    with open(OUTPUT_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(data_store, f, indent=2, ensure_ascii=False)
    
    total_weeks = sum(len(weeks) for b in data_store.values() for weeks in b.values())
    print(f"\nSuccessfully written {total_weeks} total weeks across {len(data_store)} branches into {OUTPUT_JSON_PATH}")

if __name__ == "__main__":
    parse_docx()
