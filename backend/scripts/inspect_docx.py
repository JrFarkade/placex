import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r"C:\Users\JrFar\.gemini\antigravity\scratch\placex\backend\PlaceX_Roadmap_Corrected_Reference.docx"
doc = docx.Document(docx_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i} [{p.style.name}]: {p.text.strip()}")

print(f"\nTotal Tables: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    header = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells]
    num_rows = len(table.rows)
    print(f"\n--- Table {t_idx + 1} ({num_rows} rows, {len(header)} cols) ---")
    print(f"Headers: {header}")
    if num_rows > 1:
        first_row = [c.text.strip().replace('\n', ' ') for c in table.rows[1].cells]
        print(f"Row 1 Sample: {first_row[:3]}")
