import docx

def update_methodology_document():
    doc = docx.Document('METHODOLOGY.docx')
    paragraphs = doc.paragraphs
    tables = doc.tables

    # Paragraph text updates
    # P6: Introduction
    paragraphs[6].text = (
        "The methodology defines the structured engineering approach adopted to design, develop, and implement "
        "PlaceX, an integrated Artificial Intelligence-powered career operating system for undergraduate engineering "
        "students. The overarching objective of this methodology is to build a high-performance, modular, and scalable "
        "platform that unifies fragmented placement preparation activities into a cohesive, data-driven workflow. "
        "A hybrid Software Development Life Cycle (SDLC) model was selected, combining structured architectural planning "
        "with iterative, component-level refinement. This engineering methodology balances rigorous system boundaries with "
        "the agility needed to tune generative reasoning prompts, deterministic scoring algorithms, and multi-modal interaction "
        "loops based on empirical evaluations. The methodology spans six distinct development phases, ensuring end-to-end "
        "traceability from initial problem diagnosis through operational deployment."
    )

    # P13: Overview of Methodology
    paragraphs[13].text = (
        "The development of PlaceX is executed across six structured phases: requirement analysis, modular system design, "
        "context and data preparation, component implementation, multi-tier testing and evaluation, and deployment with "
        "operational maintenance. Each phase governs a vital dimension of system dependability and performance. The system "
        "workflow captures student telemetry across specialized preparation modules, synchronizes analytical records within "
        "a persistent database, and exposes aggregated evaluation vectors to the central Host Agent orchestration engine to deliver "
        "actionable, real-time mentorship guidance."
    )

    # P16: Phase 1: Requirement Analysis
    paragraphs[16].text = (
        "The initial phase focuses on diagnosing undergraduate placement preparation challenges and formulating technical system "
        "specifications. Students traditionally suffer from fragmented preparation ecosystems, switching between disconnected tools "
        "for resume formatting, coding practice, mock interviews, and roadmap tracking without unified skill-gap feedback. "
        "Functional requirements encompass authenticated student access, ATS-compliant resume parsing and job description matching, "
        "in-browser multi-language code execution with abstract syntax tree (AST) complexity analysis, live multi-modal AI mock "
        "interviews, adaptive knowledge testing, automated career roadmapping, real-time notification dispatch, and a centralized "
        "Host Agent. Non-functional requirements include sub-second local code evaluation, robust exception handling, strict "
        "database integrity via foreign key constraints, high usability across responsive screen viewports, and secure credential storage."
    )

    # P19: Phase 2: System Design
    paragraphs[19].text = (
        "In this phase, empirical requirements are translated into a scalable, multi-layered architecture. PlaceX follows a client-server "
        "paradigm featuring a modern single-page frontend application, a high-throughput asynchronous REST backend, an AI reasoning engine, "
        "and relational data persistence. The presentation tier is built in React 18 and TypeScript with Tailwind CSS, delivering responsive "
        "workspaces for code editing, real-time video simulation, and analytical telemetry. The application tier is powered by Python and "
        "FastAPI, exposing structured RESTful endpoints organized by domain routers. Relational persistence is managed through SQLite and "
        "SQLAlchemy ORM. The system architecture enforces a clean separation of concerns: specialized operational engines independently "
        "compute deterministic metrics (such as ATS keyword density, AST execution cyclomatic depth, and test suite pass rates), while the "
        "central Host Agent orchestrator ingests normalized context vectors to formulate strategic advice via Google Gemini."
    )

    # P21: Section Heading 3.2.3
    paragraphs[21].text = "3.2.3 Phase 3: Data and Context Preparation "

    # P22: Phase 3 Paragraph 1
    paragraphs[22].text = (
        "The platform ingests, processes, and maintains multi-modal career preparation data across several structured domains. "
        "Internal persistence stores student profile metadata, uploaded resumes in PDF and DOCX formats, synthesized question blueprints, "
        "code submissions, assessment attempts, and time-stamped activity events."
    )

    # P24: Phase 3 Paragraph 2
    paragraphs[24].text = (
        "Data preprocessing involves extracting raw text from uploaded resumes using PyMuPDF and python-docx, parsing structural sections, "
        "and isolating technical skill keywords against a canonical taxonomy. Job descriptions undergo TF-IDF vectorization and keyword extraction "
        "to calculate cosine similarity scores. Student coding submissions are statically parsed into Abstract Syntax Trees using Python's native ast "
        "module to extract time and space complexity indicators prior to execution."
    )

    # P27: Phase 4: Implementation
    paragraphs[27].text = (
        "The implementation phase realizes the system architecture through modular, decoupled software components. The frontend workspace "
        "integrates Monaco Editor for syntax-highlighted programming, interactive analytics charts, and a slide-over Host Agent drawer. "
        "The backend implements dedicated micro-services for resume analysis, coding evaluation, interview synthesis, knowledge testing, and "
        "readiness aggregation. Coding submissions are validated through a hybrid execution engine that leverages Judge0 API alongside an isolated "
        "local Python subprocess runner equipped with base64 transport, linecache injection, and memory monitoring. The mock interview engine "
        "researches target companies via Tavily web intelligence, generates tiered behavioral and technical questions via Gemini, and coordinates "
        "real-time speech recognition and video telemetry. The Host Agent operates as a central cognitive copilot, synthesizing student state across "
        "all modules into actionable next-step interventions."
    )

    # P30: Phase 5: Testing and Evaluation
    paragraphs[30].text = (
        "PlaceX undergoes rigorous multi-tier testing to validate functional correctness, execution safety, and user experience integrity. "
        "Unit testing verifies discrete utility functions including resume section segmenters, AST tree walkers, and quiz scoring routines. "
        "Integration testing validates synchronous communication between FastAPI REST endpoints, SQLAlchemy database sessions, and the React client. "
        "System testing exercises complete end-to-end workflows, verifying that actions taken in the coding sandbox or mock interview room reliably update "
        "the student's placement readiness score and trigger contextual Host Agent notifications. Performance and usability testing confirm smooth "
        "editor responsiveness, low-latency API dispatch, and intuitive cross-device interface navigation."
    )

    # P36: Phase 6: Deployment and Maintenance
    paragraphs[36].text = (
        "PlaceX is deployed using standard modern runtime environments, utilizing Uvicorn as an asynchronous ASGI server for FastAPI and Vite for "
        "optimized production asset bundling. Relational persistence is maintained through ACID-compliant SQLite databases with structured SQLAlchemy "
        "migrations. Configuration parameters, security secrets, and API access tokens are securely managed through isolated environment variables. "
        "Continuous maintenance protocols encompass API regression monitoring, dependency security auditing, prompt template refinement, and test case "
        "expansion to ensure sustained platform reliability and student preparation fidelity."
    )

    # P39: Advantages of Methodology
    paragraphs[39].text = (
        "The adopted methodology yields significant architectural and operational benefits for AI-driven educational software. "
        "Its modular architecture isolates domain failures and facilitates independent component scaling. The deliberate architectural separation "
        "between deterministic algorithmic computation (ATS parsing, AST analysis, test validation) and generative LLM reasoning eliminates AI "
        "hallucinations in critical student assessments while preserving high conversational empathy in mentorship. Centralizing student telemetry "
        "within the Host Agent guarantees coherent, longitudinal guidance across preparation domains, offering undergraduate students an integrated, "
        "rigorous, and actionable pathway to campus placement readiness."
    )

    # P42: System Workflow
    paragraphs[42].text = (
        "The operational workflow of PlaceX begins when an authenticated student accesses the platform to pursue targeted placement preparation. "
        "Upon authenticating and establishing their target job role, company profile, and seniority level, the student navigates through specialized "
        "career modules. Each module captures domain-specific artifacts, computes deterministic performance metrics, persists longitudinal progress records, "
        "and emits standardized telemetry to the central Host Agent orchestration engine. The Host Agent analyzes the aggregated preparation state to synthesize "
        "actionable recommendations, update global readiness indices, and guide ongoing student development."
    )

    # P43: Authentication & Onboarding Workspace
    paragraphs[43].runs[0].text = "Authentication and Onboarding: "
    paragraphs[43].runs[1].text = (
        "The authentication interface serves as the entry portal for PlaceX, providing students with secure credentials management. "
        "It supports standard email registration alongside Google OAuth 2.0 single sign-on. During initial onboarding, students define their academic "
        "background, target software engineering roles, dream companies, and primary technical domains, establishing the baseline profile necessary "
        "to personalize all downstream preparation modules and analytical rubrics."
    )

    # P44: Student Profile & External Integrations
    paragraphs[44].runs[0].text = "Student Profile and External Sync: "
    paragraphs[44].runs[1].text = (
        "The profile workspace enables students to maintain comprehensive academic and professional identities. Beyond standard biographical details, "
        "it integrates external verified platforms including GitHub repositories, LeetCode contest statistics, and LinkedIn credentials. This module "
        "validates external handles, fetches live problem-solving tallies, and incorporates external achievements directly into the platform's career database."
    )

    # P45: ATS Resume Analyzer & JD Matcher
    paragraphs[45].runs[0].text = "ATS Resume Analyzer and Matcher: "
    paragraphs[45].runs[1].text = (
        "This module allows students to upload resumes in PDF or DOCX formats for automated evaluation against industry standards. It executes a dual-mode "
        "analysis: Mode A computes a comprehensive ATS health score across contact formatting, section completeness, action verbs, and quantifiable impact, "
        "while Mode B evaluates keyword alignment against specific job descriptions using TF-IDF vectorization and semantic matching to identify skill gaps."
    )

    # P46: Interactive Coding Sandbox
    paragraphs[46].runs[0].text = "Interactive Coding Sandbox: "
    paragraphs[46].runs[1].text = (
        "The coding sandbox provides an in-browser development environment powered by Monaco Editor. Students solve algorithmic problems across multiple "
        "programming languages with custom test cases, automated multi-case grading, and static AST complexity analysis. The sandbox features an integrated "
        "Host Agent Explain and Fix assistant that analyzes compilation errors and algorithmic bugs to provide structured, Socratic debugging hints."
    )

    # P47: AI Mock Interview Subsystem
    paragraphs[47].runs[0].text = "AI Mock Interview Module: "
    paragraphs[47].runs[1].text = (
        "The mock interview module simulates realistic placement interviews through a multi-stage process. The Prep Brain conducts automated company intelligence "
        "research via Tavily and synthesizes tiered technical and behavioral question banks. During live simulation, real-time speech-to-text captures student responses, "
        "audio prosody and facial geometry evaluate non-verbal communication, and each turn is scored against calibrated rubrics before generating a comprehensive debrief."
    )
    # Clear remaining stale runs in P47
    for r in paragraphs[47].runs[2:]:
        r.text = ""

    # P48: Adaptive Roadmap & Knowledge Quiz
    paragraphs[48].runs[0].text = "Adaptive Roadmap and Knowledge Quiz: "
    paragraphs[48].runs[1].text = (
        "This combined module guides structured learning and conceptual verification. The Career Roadmap dynamically sequences milestone topics, recommended resources, "
        "and practical milestones tailored to the student's target role. The Knowledge Quiz module tests foundational computer science domains, evaluating accuracy and "
        "speed across randomized question banks to identify specific conceptual weaknesses."
    )

    # P49: Mission Control Dashboard & Host Agent
    paragraphs[49].runs[0].text = "Mission Control Dashboard and Host Agent: "
    paragraphs[49].runs[1].text = (
        "The PlaceX dashboard functions as the centralized mission control, consolidating metrics from all preparatory modules into an interactive command interface. "
        "It visualizes overall placement readiness percentages, weekly activity streaks, competency radar breakdowns, and real-time notification alerts. "
        "Operating ubiquitously across the dashboard, the Main Host Agent maintains multi-turn conversational context, analyzes student performance history, "
        "and proactively recommends the highest-leverage preparation tasks to maximize campus recruitment success."
    )
    # Clear remaining stale runs in P49
    for r in paragraphs[49].runs[2:]:
        r.text = ""

    # P66: Tools and Technologies introductory text
    paragraphs[66].text = (
        "The PlaceX career operating system leverages a modern, robust full-stack technology ecosystem. The presentation layer is built with React.js 18, "
        "TypeScript, Vite, and Tailwind CSS, incorporating Monaco Editor for programming workflows. The backend application layer is developed with Python 3.12 "
        "and FastAPI, providing asynchronous RESTful APIs and robust data validation. Relational data persistence is managed via SQLAlchemy ORM and SQLite. "
        "Cognitive reasoning and conversational orchestration are powered by Google Gemini, while Tavily Search API, Deepgram speech recognition, and Judge0 "
        "deliver specialized company intelligence, multi-modal interview analysis, and secure code execution."
    )

    # P75: System Architecture overview
    paragraphs[75].text = (
        "PlaceX is engineered upon a decoupled, five-tier modular architecture designed to maximize operational maintainability, horizontal component extensibility, "
        "and low-latency user interactivity. The architecture cleanly segregates presentation delivery, application business logic, artificial intelligence reasoning, "
        "relational data persistence, and specialized external execution services."
    )

    # P78: 3.6.1 Presentation Layer
    paragraphs[78].text = (
        "The Presentation Layer comprises the client-side single-page application built using React 18, TypeScript, Vite, and Tailwind CSS. It delivers responsive, "
        "high-contrast workspaces for the central dashboard, ATS resume diagnostics, coding sandbox IDE with Monaco Editor, live video interview rooms, skill roadmaps, "
        "and the slide-over Host Agent workspace. It handles local state management, optimistic UI updates, and secure JWT authentication storage."
    )

    # P81: 3.6.2 Application Layer
    paragraphs[81].text = (
        "The Application Layer represents the core backend server developed using Python and FastAPI. It exposes structured RESTful API routers governing authentication, "
        "resume evaluation, code execution coordination, interview blueprint synthesis, quiz assessment, profile management, notifications, and Host Agent reasoning. "
        "This layer handles JWT authorization, request payload validation via Pydantic schemas, and orchestration of domain services."
    )

    # P86: 3.6.3 AI Processing Layer
    paragraphs[86].text = (
        "The AI Processing Layer serves as the intelligence core of PlaceX. Powered by Google Gemini large language models, it hosts the Main Host Agent reasoning engine, "
        "dynamic resume improvement generators, Socratic coding error diagnostic assistants, and calibrated interview question synthesizers. The layer operates on structured "
        "context prompts, enforcing deterministic output formatting and conversational mentor personas without contaminating algorithmic scoring routines."
    )

    # P89: 3.6.4 Data Layer
    paragraphs[89].text = (
        "The Data Layer guarantees ACID-compliant persistence and relational integrity across the platform using SQLite and SQLAlchemy ORM. It securely stores student account "
        "records, detailed academic profiles, parsed resume data, job descriptions, coding problem catalogues and submission histories, interview question rubrics and evaluations, "
        "quiz results, connected profile tokens, and time-stamped system notifications."
    )

    # P92: 3.6.5 Integration Layer
    paragraphs[92].text = (
        "The Integration Layer connects PlaceX with specialized external engines, third-party cloud services, and verified developer platforms. For secure code execution, "
        "it interfaces with the Judge0 API alongside an isolated local Python execution runner to grade code against multiple test inputs safely. In the mock interview module, "
        "it integrates Tavily web search API to extract up-to-date company technical stack signals and interview trends, and coordinates Deepgram automated speech recognition "
        "for real-time vocal transcription. For professional identity verification, the layer interfaces with Google OAuth 2.0, GitHub GraphQL APIs, and LeetCode public endpoints "
        "to pull verified contest badges, problem-solving numbers, and open-source contributions directly into the student's profile. This multi-system integration architecture "
        "ensures that PlaceX functions as a comprehensive, real-world career operating system."
    )

    # ==================== TABLE 0 ====================
    # Table 3.1: Overview of Methodology
    t0 = tables[0]
    t0.rows[1].cells[1].paragraphs[0].text = "Diagnosing placement preparation challenges, analyzing user needs, and establishing functional and non-functional specifications."
    t0.rows[2].cells[1].paragraphs[0].text = "Designing the multi-tier client-server architecture, relational data schemas, and Host Agent orchestration boundaries."
    # R3 C0 has 2 paragraphs: ['Phase 3: Data ', 'Preparation ']
    # Update to: ['Phase 3: Data ', 'and Context ']
    t0.rows[3].cells[0].paragraphs[1].text = "and Context "
    t0.rows[3].cells[1].paragraphs[0].text = "Extracting resume text, parsing job descriptions, defining skill taxonomies, and building AST evaluation pipelines."
    t0.rows[4].cells[1].paragraphs[0].text = "Developing the React frontend, FastAPI backend services, Monaco coding sandbox, and Gemini AI reasoning modules."
    t0.rows[5].cells[1].paragraphs[0].text = "Conducting unit tests, integration tests, system workflows, execution security verification, and usability assessments."
    t0.rows[6].cells[1].paragraphs[0].text = "Configuring runtime environments, ASGI server hosting, database migrations, and operational maintenance routines."

    # ==================== TABLE 1 ====================
    # Table 3.2: Types of Testing
    t1 = tables[1]
    t1.rows[1].cells[1].paragraphs[0].text = "Validation of individual functions: AST parsing, resume section extraction, and quiz scoring algorithms."
    t1.rows[2].cells[1].paragraphs[0].text = "Testing synchronous interaction between FastAPI REST endpoints, SQLAlchemy database sessions, and UI components."
    t1.rows[3].cells[1].paragraphs[0].text = "End-to-end verification of complete student journeys from resume upload to mock interview and readiness aggregation."
    t1.rows[4].cells[1].paragraphs[0].text = "Benchmarking API latency, code execution sandbox timeouts, and real-time Host Agent response speeds."
    t1.rows[5].cells[1].paragraphs[0].text = "Evaluating interface accessibility, responsive layout clarity, Monaco Editor ergonomics, and navigation flow."

    # ==================== TABLE 2 ====================
    # Table 3.3: Tools and Technologies
    t2 = tables[2]
    # Row 1: Frontend
    t2.rows[1].cells[0].paragraphs[0].text = "Frontend Tier "
    t2.rows[1].cells[1].paragraphs[0].text = "React.js 18, TypeScript, Vite, Tailwind CSS, Monaco Editor "
    t2.rows[1].cells[2].paragraphs[0].text = "Delivers responsive, accessible user interfaces, interactive code editing, and real-time telemetry visualization."
    if len(t2.rows[1].cells[2].paragraphs) > 1:
        t2.rows[1].cells[2].paragraphs[1].text = ""

    # Row 2: Backend
    t2.rows[2].cells[0].paragraphs[0].text = "Backend Tier "
    t2.rows[2].cells[1].paragraphs[0].text = "Python 3.12, FastAPI, Uvicorn, Pydantic, RESTful APIs "
    t2.rows[2].cells[2].paragraphs[0].text = "Executes core application logic, manages RESTful endpoints, coordinates micro-services, and authenticates requests."

    # Row 3: Database
    t2.rows[3].cells[0].paragraphs[0].text = "Data Tier "
    t2.rows[3].cells[1].paragraphs[0].text = "SQLite, SQLAlchemy ORM, Alembic Migrations "
    t2.rows[3].cells[2].paragraphs[0].text = "Stores student profiles, resumes, coding submissions, interview evaluations, quiz results, and activity telemetry."

    # Row 4: AI & Execution Services
    t2.rows[4].cells[0].paragraphs[0].text = "AI & External Services "
    t2.rows[4].cells[1].paragraphs[0].text = "Google Gemini LLM, Judge0 API, Tavily API, Deepgram STT "
    if len(t2.rows[4].cells[1].paragraphs) > 1:
        t2.rows[4].cells[1].paragraphs[1].text = ""
    if len(t2.rows[4].cells[1].paragraphs) > 2:
        t2.rows[4].cells[1].paragraphs[2].text = ""
    t2.rows[4].cells[2].paragraphs[0].text = "Drives Host Agent reasoning, automated interview question synthesis, company research, and isolated code execution."
    if len(t2.rows[4].cells[2].paragraphs) > 1:
        t2.rows[4].cells[2].paragraphs[1].text = ""

    doc.save('METHODOLOGY.docx')
    print("METHODOLOGY.docx successfully updated and saved!")

if __name__ == '__main__':
    update_methodology_document()
